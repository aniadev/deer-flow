#!/usr/bin/env python3
"""Fill {{TOKEN}} placeholders in a vn-admin-docs template.

Usage:
    python3 fill_template.py --template cong_van.docx --data data.json --out out.docx

data.json: {"SO_KY_HIEU": "[...]", "NOI_DUNG": "Đoạn 1.\nĐoạn 2.", ...}
- Formatting is preserved: each token lives in its own run in the template.
- NOI_DUNG supports multi-paragraph content: split on "\n", each new paragraph
  clones the template paragraph's format (justify, first-line indent, spacing).
- Missing keys are left as {{TOKEN}} so the audit step can report them.
"""
import argparse, copy, json, re, sys

from docx import Document

TOKEN = re.compile(r"\{\{([A-Z_]+)\}\}")


def iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def fill_runs(p, data):
    for r in p.runs:
        def sub(m):
            return str(data[m.group(1)]) if m.group(1) in data else m.group(0)
        new = TOKEN.sub(sub, r.text)
        if new != r.text:
            r.text = new


def expand_multiparagraph(doc, data, token="NOI_DUNG"):
    """If NOI_DUNG contains newlines, split into cloned paragraphs first."""
    value = data.get(token)
    if not value or "\n" not in str(value):
        return
    parts = [s.strip() for s in str(value).split("\n") if s.strip()]
    for p in iter_paragraphs(doc):
        if any("{{%s}}" % token in r.text for r in p.runs):
            anchor = p._p
            for part in reversed(parts[1:]):
                clone = copy.deepcopy(p._p)
                anchor.addnext(clone)
                # rewrite the clone's token text
                from docx.text.paragraph import Paragraph
                cp = Paragraph(clone, p._parent)
                for r in cp.runs:
                    if "{{%s}}" % token in r.text:
                        r.text = r.text.replace("{{%s}}" % token, part)
            data = dict(data)
            for r in p.runs:
                if "{{%s}}" % token in r.text:
                    r.text = r.text.replace("{{%s}}" % token, parts[0])
            return


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--data", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)

    doc = Document(args.template)
    expand_multiparagraph(doc, data)
    for p in iter_paragraphs(doc):
        fill_runs(p, data)
    doc.save(args.out)

    # Report unfilled tokens
    leftover = sorted({m for p in iter_paragraphs(Document(args.out))
                       for r in p.runs for m in TOKEN.findall(r.text)})
    print(f"Saved: {args.out}")
    if leftover:
        print("UNFILLED TOKENS (placeholders kept intentionally or missing data):")
        for tk in leftover:
            print(f"  - {{{{{tk}}}}}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
